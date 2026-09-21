# -*- coding: utf-8 -*-
"""Full independent audit of every formula cell in the delivered workbook."""
import openpyxl
import sys
SRC=sys.argv[1] if len(sys.argv)>1 else "/home/user/Health-Insurance-PEO/deliverables/Onyx_Angle_Questco_Cost_Model.xlsx"
wf=openpyxl.load_workbook(SRC)                 # formulas
wv=openpyxl.load_workbook(SRC,data_only=True)  # cached values

FAIL=[]; CHECKED=set()
def V(sh,c): return wv[sh][c].value
def chk(sh,c,exp,tol=0.005):
    CHECKED.add((sh,c)); got=V(sh,c)
    if isinstance(exp,(int,float)) and not isinstance(exp,bool):
        ok = isinstance(got,(int,float)) and abs(got-exp)<=tol
    else:
        ok = got==exp
    if not ok: FAIL.append("%s!%s: got %r  expected %r"%(sh,c,got,exp))
def skip(sh,c): CHECKED.add((sh,c))

# ============ SOURCE DATA, read only from the workbook ============
AP="Angle Plans"; CO="Contributions"; CT="Cost by Tier"; EL="Elections"
BE="By Employee"; QU="Questco"; KK="401(k)"; CB="Combined"
plans=[]; 
for r in range(5,8):
    plans.append(dict(name=V(AP,f"A{r}"),hsa=V(AP,f"B{r}"),di=V(AP,f"C{r}"),df=V(AP,f"D{r}"),
                      oi=V(AP,f"E{r}"),of=V(AP,f"F{r}"),
                      prem={"EE":V(AP,f"L{r}"),"ES":V(AP,f"M{r}"),"EC":V(AP,f"N{r}"),"FAM":V(AP,f"O{r}")}))
P={p["name"]:p for p in plans}; ORDER=[p["name"] for p in plans]
CODES=["EE","ES","EC","FAM"]
basis=V(CO,"C6"); minpct=V(CO,"C19")
tier={V(CO,f"B{r}"):dict(inc=V(CO,f"C{r}"),pct=V(CO,f"D{r}"),fix=V(CO,f"E{r}")) for r in range(11,15)}
MIN=minpct*min(p["prem"]["EE"] for p in plans)
emp=[dict(n=V(CO,f"A{r}"),pay=V(CO,f"B{r}"),tier=V(CO,f"C{r}"),code=V(CO,f"D{r}"),yn=V(CO,f"E{r}"))
     for r in range(26,31)]

# ============ CONTRIBUTIONS ============
chk(CO,"C20",MIN)
for i,code in enumerate(CODES):
    r=11+i
    chk(CO,f"F{r}",MIN)
    chk(CO,f"G{r}",min(p["prem"][code] for p in plans))
    t=tier[code]
    if t["inc"]!="Yes": st="Excluded - carried for history, not counted anywhere"
    elif basis=="Fixed dollars":
        st="Funded at the fixed amount" if t["fix"]>=MIN else "Topped up to the Angle minimum"
    else:
        st="Funded at the percentage" if t["pct"]*min(p["prem"][code] for p in plans)>=MIN \
           else "Topped up to the Angle minimum on the cheapest plan"
    chk(CO,f"H{r}",st)
tot_enr=0; tot_cnt=0
for i,code in enumerate(CODES):
    r=35+i
    enr=sum(1 for e in emp if e["code"]==code and e["yn"]=="Yes")
    cnt=enr if tier[code]["inc"]=="Yes" else 0
    chk(CO,f"C{r}",tier[code]["inc"]); chk(CO,f"D{r}",enr); chk(CO,f"E{r}",cnt)
    tot_enr+=enr; tot_cnt+=cnt
chk(CO,"D39",tot_enr); chk(CO,"E39",tot_cnt); chk(CO,"D40",tot_enr)
chk(CO,"C21",MIN*tot_enr*12)

# ============ COST BY TIER ============
ct={}; grp={}
r=5
for pl in ORDER:
    g=[]
    for code in CODES:
        prem=P[pl]["prem"][code]; t=tier[code]
        enr=0 if t["inc"]!="Yes" else sum(1 for e in emp if e["code"]==code and e["yn"]=="Yes")
        contrib=t["fix"] if basis=="Fixed dollars" else t["pct"]*prem
        onyx=min(prem,max(contrib,MIN)); eec=prem-onyx
        ded=P[pl]["di"] if code=="EE" else P[pl]["df"]
        oop=P[pl]["oi"] if code=="EE" else P[pl]["of"]
        d=dict(D=t["inc"],E=enr,F=prem,G=contrib,H=MIN,I=onyx,J=eec,
               K=(onyx/prem if prem else 0),L=onyx*enr,M=onyx*enr*12,N=eec*enr,O=eec*enr*12,
               P=eec*12,Q=ded,R=oop,S=eec*12+oop,T=f"{pl}|{code}")
        for col,v in d.items(): chk(CT,f"{col}{r}",v)
        ct[(pl,code)]=d; g.append(d); r+=1
    chk(CT,f"E{r}",sum(x["E"] for x in g))
    for col in "LMNO": chk(CT,f"{col}{r}",sum(x[col] for x in g))
    chk(CT,f"S{r}",sum(x["S"]*x["E"] for x in g))
    grp[pl]=dict(L=sum(x["L"] for x in g),M=sum(x["M"] for x in g),
                 N=sum(x["N"] for x in g),O=sum(x["O"] for x in g))
    r+=2

# ============ ELECTIONS - main table ============
el_on=[]; el_ee=[]; el_wc=[]; el_prem_counted=0.0
for i,e in enumerate(emp):
    r=7+i
    chk(EL,f"B{r}",e["yn"]); chk(EL,f"D{r}",e["code"])
    chk(EL,f"F{r}",tier[e["code"]]["inc"])
    pl=V(EL,f"E{r}"); row=ct[(pl,e["code"])]
    chk(EL,f"G{r}",row["F"])
    gate = (e["yn"]=="Yes" and tier[e["code"]]["inc"]=="Yes")
    on=row["I"] if gate else 0.0; ee=row["J"] if gate else 0.0
    chk(EL,f"H{r}",on); chk(EL,f"I{r}",ee); chk(EL,f"J{r}",on+ee)
    chk(EL,f"K{r}",on*12); chk(EL,f"L{r}",ee*12); chk(EL,f"M{r}",(on+ee)*12)
    wc = row["S"] if gate else 0.0
    chk(EL,f"N{r}",wc); el_wc.append(wc)
    el_on.append(on); el_ee.append(ee)
    if gate: el_prem_counted+=row["F"]
ON_M=sum(el_on); EE_M=sum(el_ee)
chk(EL,"B12","%d enrolled"%sum(1 for e in emp if e["yn"]=="Yes"))
for c,v in (("H12",ON_M),("I12",EE_M),("J12",ON_M+EE_M),
            ("K12",ON_M*12),("L12",EE_M*12),("M12",(ON_M+EE_M)*12),("N12",sum(el_wc))):
    chk(EL,c,v)
chk(EL,"C14",ON_M); chk(EL,"C15",ON_M*12); chk(EL,"C16",EE_M); chk(EL,"C17",EE_M*12)

# ============ ELECTIONS - saved scenarios ============
PAIRS=[("E","F"),("G","H"),("I","J"),("K","L")]
base=None
for pc,dc in PAIRS:
    onyx=0.0
    for i,e in enumerate(emp):
        r=27+i; u=7+i
        prem=V(EL,f"G{u}"); chk(EL,f"D{r}",prem)
        chk(EL,f"A{r}",e["n"]); chk(EL,f"B{r}",e["code"]); chk(EL,f"C{r}",V(EL,f"E{u}"))
        gate=(e["yn"]=="Yes" and tier[e["code"]]["inc"]=="Yes")
        v=min(prem,max(V(EL,f"{pc}{r}")*prem,MIN)) if gate else 0.0
        chk(EL,f"{dc}{r}",v); onyx+=v
    chk(EL,f"{dc}32",onyx); chk(EL,f"{dc}33",onyx*12)
    chk(EL,f"{dc}34",el_prem_counted-onyx); chk(EL,f"{dc}35",(el_prem_counted-onyx)*12)
    if base is None: base=(onyx*12,(el_prem_counted-onyx)*12)
    chk(EL,f"{dc}36",onyx*12-base[0]); chk(EL,f"{dc}37",(el_prem_counted-onyx)*12-base[1])
d=V(EL,"H36")
exp=("Scenario 2 costs Onyx exactly what Scenario 1 does." if abs(d)<1e-9 else
     "Scenario 2 costs Onyx ${:,.0f}".format(abs(d))+(" MORE" if d>0 else " LESS")+
     " a year than Scenario 1 - and moves that same amount "+("off" if d>0 else "onto")+" the employees.")
chk(EL,"A39",exp)

# ============ BY EMPLOYEE ============
r=5
for i,e in enumerate(emp):
    for pl in ORDER:
        row=ct[(pl,e["code"])]
        chk(BE,f"B{r}",e["yn"])
        for col,src in (("F","D"),("G","F"),("H","I"),("I","J"),("J","P"),("K","Q"),("L","R"),("M","S")):
            chk(BE,f"{col}{r}",row[src])
        r+=1

# ============ QUESTCO ============
admin=V(QU,"C6")*V(QU,"D6")*12
wc = 0 if V(QU,"E7")=="No" else V(QU,"C7")
cy = 0 if V(QU,"E8")=="No" else V(QU,"C8")
for c,v in (("G6",admin),("F6",admin/12),("G7",wc),("F7",wc/12),("G8",cy),("F8",cy/12),
            ("F9",(admin+wc+cy)/12),("G9",admin+wc+cy)):
    chk(QU,c,v)
anc=0.0
for r in (15,16,17):
    a = V(QU,f"D{r}")*12 if V(QU,f"B{r}")=="Yes" else 0
    chk(QU,f"E{r}",a); anc+=a
chk(QU,"D18",anc/12); chk(QU,"E18",anc)
chk(QU,"B26",sum(V(QU,f"B{r}") for r in (23,24,25)))

# ============ 401(k) ============
fee=V(KK,"C7"); paidby=V(KK,"C8"); climit=V(KK,"C9"); dlimit=V(KK,"C10")
mtype=V(KK,"C12"); t1r,t1p=V(KK,"C13"),V(KK,"D13"); t2r,t2p=V(KK,"C14"),V(KK,"D14"); nel=V(KK,"C15")
S={"B":0,"E":0,"F":0,"H":0,"I":0,"J":0}
for i in range(5):
    r=19+i
    pay=V(KK,f"B{r}"); part=V(KK,f"C{r}"); dfr=V(KK,f"D{r}")
    E = 0 if part!="Yes" else min(min(pay,climit)*dfr,dlimit)
    I = 0 if part!="Yes" else min(pay,climit)*(t1r*min(dfr,t1p)+t2r*min(max(dfr-t1p,0),t2p))
    J = 0 if part!="Yes" else min(pay,climit)*nel
    F = 0 if part!="Yes" else (0 if mtype=="None" else (J if mtype=="Non-elective" else I))
    H = 0 if part!="Yes" else fee*4
    for c,v in (("E",E),("F",F),("G",0 if pay==0 else F/pay),("H",H),("I",I),("J",J)):
        chk(KK,f"{c}{r}",v)
    S["B"]+=pay; S["E"]+=E; S["F"]+=F; S["H"]+=H; S["I"]+=I; S["J"]+=J
for c in "BEFHIJ": chk(KK,f"{c}24",S[c])
chk(KK,"C24","%d in"%sum(1 for i in range(5) if V(KK,f"C{19+i}")=="Yes"))
chk(KK,"G24",0 if S["B"]==0 else S["F"]/S["B"])
chk(KK,"B30",S["I"]); chk(KK,"B31",S["J"])
for r,b in ((29,V(KK,"B29")),(30,S["I"]),(31,S["J"])): chk(KK,f"C{r}",b/12)
c35=S["F"]; c36=S["H"] if paidby=="Employer" else 0
chk(KK,"C35",c35); chk(KK,"B35",c35/12); chk(KK,"C36",c36); chk(KK,"B36",c36/12)
chk(KK,"C37",c35+c36); chk(KK,"B37",(c35+c36)/12)
K_M=(c35+c36)/12; K_A=c35+c36

# ============ COMBINED ============
lines=[(ON_M,ON_M*12),(admin/12,admin),(wc/12,wc),(cy/12,cy),(anc/12,anc),(K_M,K_A)]
for i,(m,a) in enumerate(lines):
    chk(CB,f"C{7+i}",m); chk(CB,f"D{7+i}",a)
rec_m=sum(x[0] for x in lines); rec_a=sum(x[1] for x in lines)
chk(CB,"C13",rec_m); chk(CB,"D13",rec_a)
chk(CB,"D14",V(QU,"G10")); chk(CB,"D15",rec_a+V(QU,"G10")); chk(CB,"C15",(rec_a+V(QU,"G10"))/12)
chk(CB,"B22",EE_M); chk(CB,"C22",EE_M*12)
chk(CB,"B23",S["E"]/12); chk(CB,"C23",S["E"])
f24m = 0 if paidby=="Employer" else S["H"]/12
f24a = 0 if paidby=="Employer" else S["H"]
chk(CB,"B24",f24m); chk(CB,"C24",f24a)
chk(CB,"B25",ON_M+EE_M); chk(CB,"C25",(ON_M+EE_M)*12)
chk(CB,"B26",rec_m+EE_M+f24m); chk(CB,"C26",rec_a+EE_M*12+f24a)
chk(CB,"C27",rec_a+sum(el_wc))   # each person on the plan they elected
nonmed=admin+wc+cy+anc+K_A
for i,pl in enumerate(ORDER):
    r=31+i
    chk(CB,f"C{r}",grp[pl]["M"]); chk(CB,f"E{r}",grp[pl]["O"])
    chk(CB,f"D{r}",grp[pl]["M"]+nonmed)
    chk(CB,f"F{r}",grp[pl]["M"]+nonmed+grp[pl]["O"]+(0 if paidby=="Employer" else S["H"]))
    chk(CB,f"G{r}",grp[pl]["M"]+nonmed+sum(ct[(pl,cd)]["S"]*ct[(pl,cd)]["E"] for cd in CODES))

# ============ COVERAGE: did we miss any formula cell? ============
allf=set(); inputs=[]
for sh in wf.sheetnames:
    for row in wf[sh].iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str) and v.startswith("="): allf.add((sh,c.coordinate))
            elif v is not None and not isinstance(v,str): inputs.append((sh,c.coordinate,v))
missed=sorted(allf-CHECKED)
print("formula cells in workbook : %d"%len(allf))
print("independently recomputed  : %d"%len(allf&CHECKED))
print("NOT verified              : %d"%len(missed))
if missed:
    for sh,c in missed[:60]: print("   unverified %s!%s  =%s"%(sh,c,wf[sh][c].value))
print()
print("MISMATCHES: %d"%len(FAIL))
for f in FAIL[:60]: print("   "+f)
print()
print("hardcoded numeric inputs  : %d  (listed separately for source review)"%len(inputs))
